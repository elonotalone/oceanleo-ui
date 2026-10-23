"""Generate one small sample file per editor type into harness/fixtures/.

Run: python3 -I harness/make-fixtures.py   (idempotent; overwrites)
Uploading these through My library → Upload file gives the acceptance account
one item of each uploadable type. Rename a copy to "[Wxx] …" before uploading
so other owners can tell whose item it is.
"""
import json
import math
import struct
import subprocess
from pathlib import Path

OUT = Path(__file__).resolve().parent / "fixtures"
OUT.mkdir(exist_ok=True)


def pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = OUT / "sample.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    for page in range(1, 4):
        c.setFont("Helvetica-Bold", 20)
        c.drawString(72, h - 90, f"Acceptance form - page {page}")
        c.setFont("Helvetica", 12)
        y = h - 140
        for line in [
            "Name: ______________________",
            "Date: ______________________",
            "[ ] I agree to the terms      [ ] I do not agree",
            "This paragraph exists so text can be highlighted, underlined and",
            "struck through by the basic PDF editor during acceptance tests.",
        ]:
            c.drawString(72, y, line)
            y -= 28
        c.rect(72, 160, w - 144, 120)
        c.drawString(80, 260, "Signature box")
        c.showPage()
    c.save()


def docx():
    import docx as d

    doc = d.Document()
    doc.add_heading("Acceptance document", 1)
    doc.add_paragraph("First paragraph for text formatting tests.")
    doc.add_paragraph("Second paragraph: bold, italic, colors, lists.")
    doc.add_paragraph("Bullet one", style="List Bullet")
    doc.add_paragraph("Bullet two", style="List Bullet")
    doc.save(OUT / "sample.docx")


def text_files():
    (OUT / "sample.txt").write_text(
        "Acceptance notes\n\nThe first lines of a text file should show on the library card.\n"
        "Second line of content.\nThird line of content.\n",
        encoding="utf-8",
    )
    (OUT / "sample.md").write_text(
        "# 验收清单\n\n- 第一项：标题与开头几行要直接显示在卡片上\n- 第二项：不转圈\n\n正文第一段。\n",
        encoding="utf-8",
    )


def xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(["Task", "Owner", "Due", "Done"])
    for i in range(1, 11):
        ws.append([f"Item {i}", f"Person {i % 3}", f"2026-11-{i:02d}", i % 2 == 0])
    wb.save(OUT / "sample.xlsx")


def pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for i in range(1, 4):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Acceptance slide {i}"
        body = slide.placeholders[1].text_frame
        body.text = "First bullet"
        body.add_paragraph().text = "Second bullet"
        box = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(6), Inches(1))
        box.text_frame.text = "Free text box"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    prs.save(OUT / "sample.pptx")


def media():
    run = lambda *args: subprocess.run(args, check=True, capture_output=True)
    run("ffmpeg", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=6",
        "-ac", "2", "-b:a", "96k", str(OUT / "sample.mp3"))
    run("ffmpeg", "-y", "-f", "lavfi", "-i", "testsrc2=size=640x360:rate=25:duration=6",
        "-f", "lavfi", "-i", "sine=frequency=330:duration=6",
        "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest",
        str(OUT / "sample.mp4"))
    run("ffmpeg", "-y", "-f", "lavfi", "-i",
        "testsrc2=size=1200x800:rate=1:duration=1", "-frames:v", "1",
        str(OUT / "sample.png"))


def glb():
    # Unit cube: 24 vertices (flat normals), 36 indices.
    faces = [
        ((1, 0, 0), [(1, -1, -1), (1, 1, -1), (1, 1, 1), (1, -1, 1)]),
        ((-1, 0, 0), [(-1, -1, 1), (-1, 1, 1), (-1, 1, -1), (-1, -1, -1)]),
        ((0, 1, 0), [(-1, 1, -1), (-1, 1, 1), (1, 1, 1), (1, 1, -1)]),
        ((0, -1, 0), [(-1, -1, 1), (-1, -1, -1), (1, -1, -1), (1, -1, 1)]),
        ((0, 0, 1), [(-1, -1, 1), (1, -1, 1), (1, 1, 1), (-1, 1, 1)]),
        ((0, 0, -1), [(1, -1, -1), (-1, -1, -1), (-1, 1, -1), (1, 1, -1)]),
    ]
    positions, normals, indices = [], [], []
    for normal, quad in faces:
        base = len(positions)
        for v in quad:
            positions.append(tuple(c * 0.5 for c in v))
            normals.append(normal)
        indices += [base, base + 1, base + 2, base, base + 2, base + 3]
    pos_bytes = b"".join(struct.pack("<3f", *p) for p in positions)
    nor_bytes = b"".join(struct.pack("<3f", *n) for n in normals)
    idx_bytes = b"".join(struct.pack("<H", i) for i in indices)
    idx_bytes += b"\x00" * ((4 - len(idx_bytes) % 4) % 4)
    binary = pos_bytes + nor_bytes + idx_bytes
    gltf = {
        "asset": {"version": "2.0", "generator": "editor-overhaul fixtures"},
        "scene": 0,
        "scenes": [{"nodes": [0]}],
        "nodes": [{"mesh": 0, "name": "Cube"}],
        "meshes": [{"name": "Cube", "primitives": [{
            "attributes": {"POSITION": 0, "NORMAL": 1}, "indices": 2, "material": 0}]}],
        "materials": [{"name": "Orange", "pbrMetallicRoughness": {
            "baseColorFactor": [0.95, 0.5, 0.2, 1], "metallicFactor": 0, "roughnessFactor": 0.6}}],
        "buffers": [{"byteLength": len(binary)}],
        "bufferViews": [
            {"buffer": 0, "byteOffset": 0, "byteLength": len(pos_bytes), "target": 34962},
            {"buffer": 0, "byteOffset": len(pos_bytes), "byteLength": len(nor_bytes), "target": 34962},
            {"buffer": 0, "byteOffset": len(pos_bytes) + len(nor_bytes), "byteLength": len(indices) * 2, "target": 34963},
        ],
        "accessors": [
            {"bufferView": 0, "componentType": 5126, "count": len(positions), "type": "VEC3",
             "min": [-0.5, -0.5, -0.5], "max": [0.5, 0.5, 0.5]},
            {"bufferView": 1, "componentType": 5126, "count": len(normals), "type": "VEC3"},
            {"bufferView": 2, "componentType": 5123, "count": len(indices), "type": "SCALAR"},
        ],
    }
    js = json.dumps(gltf, separators=(",", ":")).encode()
    js += b" " * ((4 - len(js) % 4) % 4)
    total = 12 + 8 + len(js) + 8 + len(binary)
    out = struct.pack("<4sII", b"glTF", 2, total)
    out += struct.pack("<I4s", len(js), b"JSON") + js
    out += struct.pack("<I4s", len(binary), b"BIN\x00") + binary
    (OUT / "sample.glb").write_bytes(out)


if __name__ == "__main__":
    pdf()
    docx()
    text_files()
    xlsx()
    pptx()
    media()
    glb()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name}\t{f.stat().st_size}")
